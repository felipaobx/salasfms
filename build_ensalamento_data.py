import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl


TIME_PATTERN = re.compile(r"(\d{2}:\d{2})\s*-\s*(\d{2}:\d{2})")
TARGET_ROOMS = [
    "SALA 01", "SALA 02", "SALA 03", "SALA 04",
    "SALA METODOLOGIA ATIVA 5", "SALA METODOLOGIA ATIVA 6",
    "SALA 07", "SALA 8", "AUDITÓRIO",
    "TUTORIA 1", "TUTORIA 2", "TUTORIA 3", "TUTORIA 4",
]


def clean_room_name(value):
    name = re.sub(r"\s+", " ", str(value or "")).strip().upper()
    replacements = {
        "AUTIDÓRIO": "AUDITÓRIO",
        "SALA METODOLO ATIVA 5": "SALA METODOLOGIA ATIVA 5",
        "SALA METODOLO ATIVA 6": "SALA METODOLOGIA ATIVA 6",
        "LAB. HABILIDADES CLÍN/CIRÚR.": "LAB. HABILIDADES CLÍNICAS/CIRÚRGICAS",
        "LAB. PRAT.INTEGRADAS I": "LAB. PRÁTICAS INTEGRADAS I",
        "LAB. PRAT.INTEGRADAS II": "LAB. PRÁTICAS INTEGRADAS II",
        "LAB. SIMUL. DEBRIEFING 1": "LAB. SIMULAÇÃO DEBRIEFING 1",
        "LAB. SIMUL. DEBRIEFING 2": "LAB. SIMULAÇÃO DEBRIEFING 2",
    }
    return replacements.get(name, name)


def display_room_name(name):
    fixed = {
        "SALA 01": "Sala 1",
        "SALA 02": "Sala 2",
        "SALA 03": "Sala 3",
        "SALA 04": "Sala 4",
        "SALA 07": "Sala 7",
        "SALA 8": "Sala 8",
        "SALA METODOLOGIA ATIVA 5": "Sala Metodologia Ativa 5",
        "SALA METODOLOGIA ATIVA 6": "Sala Metodologia Ativa 6",
        "AUDITÓRIO": "Auditório",
        "ACOLHIMENTO": "Acolhimento",
        "MANDICAST DO SERTÃO": "Mandicast do Sertão",
        "SALA VIDEOCONFERÊNCIA": "Sala de Videoconferência",
        "LAB. ANATOMIA": "Lab. Anatomia",
        "LAB. HABILIDADES CLÍNICAS/CIRÚRGICAS": "Lab. Habilidades Clínicas/Cirúrgicas",
        "LAB. PRÁTICAS INTEGRADAS I": "Lab. Práticas Integradas I",
        "LAB. PRÁTICAS INTEGRADAS II": "Lab. Práticas Integradas II",
        "LAB. SIMULAÇÃO DEBRIEFING 1": "Lab. Simulação Debriefing 1",
        "LAB. SIMULAÇÃO DEBRIEFING 2": "Lab. Simulação Debriefing 2",
        "REPOUSO ALUNOS": "Repouso dos Alunos",
        "BIBLIOTECA": "Biblioteca",
        "CLÍNICA ESCOLA": "Clínica Escola",
    }
    if name in fixed:
        return fixed[name]
    if name.startswith("TUTORIA "):
        return "Tutoria " + name.split()[-1]
    return name.title()


def clean_event(value):
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    text = re.sub(r"(\d)�", r"\1º", text)
    replacements = {
        "PER�ODO": "PERÍODO",
        "INGL�S": "INGLÊS",
        "PR�T": "PRÁT",
        "CL�N": "CLÍN",
        "CIR�R": "CIRÚR",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


source_path = Path(sys.argv[1])
output_path = Path(sys.argv[2])
workbook = openpyxl.load_workbook(source_path, data_only=True, read_only=False)

room_names = list(TARGET_ROOMS)
room_ids = {name: index + 1 for index, name in enumerate(room_names)}
events = []

for worksheet in workbook.worksheets:
    header_rows = [
        row for row in range(1, worksheet.max_row + 1)
        if isinstance(worksheet.cell(row, 2).value, str)
        and "PER" in worksheet.cell(row, 2).value.upper()
        and worksheet.cell(row, 3).value
    ]
    for header_index, header_row in enumerate(header_rows):
        next_header = header_rows[header_index + 1] if header_index + 1 < len(header_rows) else worksheet.max_row + 1
        date_values = [
            worksheet.cell(row, 1).value
            for row in range(header_row, next_header)
            if isinstance(worksheet.cell(row, 1).value, (date, datetime))
        ]
        if not date_values:
            continue
        booking_date = date_values[0].date() if isinstance(date_values[0], datetime) else date_values[0]

        for column in range(3, worksheet.max_column + 1):
            room_key = clean_room_name(worksheet.cell(header_row, column).value)
            if room_key not in room_ids:
                continue

            daily_events = []
            for row in range(header_row + 1, next_header):
                time_value = worksheet.cell(row, 2).value
                match = TIME_PATTERN.search(str(time_value or ""))
                event_value = worksheet.cell(row, column).value
                if not match or not event_value:
                    continue
                description = clean_event(event_value)
                if not description or "INTERVALO" in description.upper() or "ALMOÇO" in description.upper():
                    continue
                start, end = match.groups()
                if daily_events and daily_events[-1]["end"] == start and daily_events[-1]["name"] == description:
                    daily_events[-1]["end"] = end
                else:
                    daily_events.append({"start": start, "end": end, "name": description})

            for item in daily_events:
                events.append({
                    "roomId": room_ids[room_key],
                    "date": booking_date.isoformat(),
                    "start": item["start"],
                    "end": item["end"],
                    "name": item["name"],
                    "ra": "",
                    "email": "",
                    "origin": "institutional",
                    "source": worksheet.title,
                })

rooms = [
    {"id": room_ids[name], "name": display_room_name(name), "capacity": None, "status": "active"}
    for name in room_names
]

events.sort(key=lambda item: (item["date"], item["start"], item["roomId"], item["name"]))
for index, event in enumerate(events, start=1):
    event["id"] = f"ensalamento-{index:05d}"

payload = (
    "// Gerado a partir de Ensalamento 2.2026.xlsx.\n"
    f"window.ENSALAMENTO_ROOMS = {json.dumps(rooms, ensure_ascii=False, separators=(',', ':'))};\n"
    f"window.ENSALAMENTO_RESERVATIONS = {json.dumps(events, ensure_ascii=False, separators=(',', ':'))};\n"
)
output_path.write_text(payload, encoding="utf-8")
print(json.dumps({
    "rooms": len(rooms),
    "reservations": len(events),
    "first_date": events[0]["date"] if events else None,
    "last_date": events[-1]["date"] if events else None,
    "by_sheet": {sheet: sum(1 for event in events if event["source"] == sheet) for sheet in workbook.sheetnames},
}, ensure_ascii=False, indent=2))
